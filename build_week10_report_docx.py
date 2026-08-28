"""Build the Week 10 DOCX report."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DERIVED = ROOT / "data" / "derived"
DATABASE = ROOT / "database"
REPORT = ROOT / "report"
OUTPUT = REPORT / "霍泓锟_第十周任务报告.docx"
RUN_DATE = "2026年8月26日"


def read_csv(name: str) -> pd.DataFrame:
    return pd.read_csv(DERIVED / name, dtype=str, keep_default_na=False)


def read_database_csv(name: str) -> pd.DataFrame:
    path = DATABASE / name
    if not path.exists():
        return pd.DataFrame()
    return pd.read_csv(path, dtype=str, keep_default_na=False)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: object, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("" if text is None else str(text))
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8.5)
    run.font.bold = bold


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D0D7DE")


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, header, bold=True)
        set_cell_shading(cell, "E8EEF5")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = Cm(widths[idx])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cells[idx].width = Cm(widths[idx])
    set_table_borders(table)
    doc.add_paragraph()
    return table


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.color.rgb = RGBColor(31, 77, 120)


def add_para(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)


def display_status(value: str) -> str:
    return {
        "READY_NEED_PASSWORD": "需密码后导入",
        "READY_WITH_PASSWORD": "可执行导入",
    }.get(str(value), str(value))


def setup_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    return doc


def main() -> None:
    REPORT.mkdir(parents=True, exist_ok=True)
    panel = read_csv("company_research_panel_week10.csv")
    board = read_csv("board_stats_week10.csv")
    desc = read_csv("descriptive_stats_week10.csv")
    fund_queue = read_csv("fund_enrichment_queue_week10.csv")
    reg = read_csv("regression_results_week10.csv")
    validation = pd.read_csv(ROOT / "validation" / "week10_validation_summary.csv", dtype=str, keep_default_na=False)
    postgres = pd.read_csv(DATABASE / "postgres_connection_audit_week10.csv", dtype=str, keep_default_na=False)
    pg_summary = read_database_csv("postgresql_disclosure_summary_week10.csv")
    pg_counts = read_database_csv("postgresql_table_counts_week10.csv")
    pg_company = read_database_csv("postgresql_company_panel_disclosure_week10.csv")
    pg_investor_type = read_database_csv("postgresql_investor_type_summary_week10.csv")
    pg_fund_status = read_database_csv("postgresql_fund_enrichment_status_week10.csv")

    doc = setup_doc()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("第十周任务报告：研究型数据集、探索性统计与数据库导入规范化")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(23, 54, 93)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(f"姓名：霍泓锟    日期：{RUN_DATE}")
    meta_run.font.name = "Microsoft YaHei"
    meta_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    meta_run.font.size = Pt(10.5)

    add_heading(doc, "一、本周任务定位", 1)
    add_para(
        doc,
        "第十周承接此前未来计划和第九周任务结果，把8家公司三表数据从“可审计、可入库准备”进一步推进为“可提出研究问题、可做探索性统计、可继续扩样”的研究型数据集。本周没有追求新增PDF数量，而是围绕变量构造、投资主体画像、基金补充核验队列、描述统计和探索性OLS展开。",
    )
    add_para(
        doc,
        "本周继续坚持“PDF未披露就留空”的工程原则。备案编码、GP和LP结构没有从三表直接披露出来时，不用主体名称推断，而是写入人工补充核验队列，等待后续回到PDF章节或使用外部权威来源核验。",
    )

    add_heading(doc, "二、核心交付结果", 1)
    summary_rows = [
        ["样本公司", str(panel["stock_code"].nunique()), "沿用统一8家公司样本"],
        ["公司层研究面板", str(len(panel)), "一家公司一行，含PE/VC强度与股权分散度指标"],
        ["投资主体画像", str(len(read_csv("investor_profile_week10.csv"))), "按主体名称和类型聚合"],
        ["基金补充核验队列", str(len(fund_queue)), "备案编码、GP、LP结构留空待核验"],
        ["探索性OLS模型", str(len(reg)), "只用于演示研究路径"],
        [
            "PostgreSQL状态",
            "临时库导入查询完成"
            if not pg_summary.empty
            else display_status(postgres.loc[postgres["check_name"].eq("postgres_import"), "status"].iloc[0]),
            "55432端口临时实例真实导入；5432主库仍需口令"
            if not pg_summary.empty
            else "需本机口令后真实导入",
        ],
    ]
    add_table(doc, ["指标", "数值", "说明"], summary_rows, [3.0, 3.0, 9.0])

    add_heading(doc, "三、公司层研究变量", 1)
    panel_rows = panel[
        [
            "stock_code",
            "company_short",
            "market",
            "broad_pevc_record_share",
            "pevc_strength_index",
            "top1_ratio_pct",
            "effective_shareholder_count",
            "quality_gate",
        ]
    ].values.tolist()
    add_table(
        doc,
        ["代码", "公司", "板块", "PE/VC占比", "强度指数", "第一大股东", "有效股东数", "质量"],
        panel_rows,
        [1.6, 2.1, 1.4, 1.8, 1.8, 1.8, 1.8, 1.5],
    )

    add_heading(doc, "四、描述性统计和板块观察", 1)
    add_para(
        doc,
        "第十周的描述性统计以公司层变量为单位。由于当前每个板块只有2家公司，以下结果只能用于形成研究假设，不能作为正式结论。它的价值主要在于验证变量能否稳定从三表生成，以及异常样本能否被单独标记。",
    )
    board_rows = board[
        [
            "market",
            "company_count",
            "avg_pevc_strength_index",
            "avg_broad_pevc_record_share",
            "avg_top1_ratio_pct",
            "avg_effective_shareholder_count",
        ]
    ].values.tolist()
    add_table(
        doc,
        ["板块", "公司数", "平均PE/VC强度", "平均PE/VC占比", "平均第一大股东", "平均有效股东数"],
        board_rows,
        [1.7, 1.5, 2.7, 2.7, 2.7, 2.7],
    )

    add_heading(doc, "五、探索性OLS结果", 1)
    reg_rows = reg[
        [
            "sample_scope",
            "dependent_variable",
            "independent_variable",
            "n",
            "coef_x",
            "r_squared",
        ]
    ].head(8).values.tolist()
    add_table(
        doc,
        ["样本", "被解释变量", "解释变量", "n", "系数", "R方"],
        reg_rows,
        [1.9, 3.6, 3.6, 1.0, 1.8, 1.4],
    )
    add_para(
        doc,
        "这些模型只用于说明研究问题如何从“论文想法”落到“变量—数据—模型”。当前样本量太小，且存在披露口径差异，因此不能把系数解释为因果关系。下一步需要扩样，并加入行业、年份、发行规模等控制变量。",
    )

    add_heading(doc, "六、基金补充核验队列", 1)
    fund_rows = fund_queue[
        ["investor_name", "investor_type_final", "manual_priority", "companies", "disclosure_status"]
    ].head(10).values.tolist()
    add_table(
        doc,
        ["投资主体", "类型", "优先级", "涉及公司", "披露状态"],
        fund_rows,
        [5.0, 1.6, 3.0, 2.0, 4.0],
    )
    add_para(
        doc,
        "本周没有把基金备案编码、GP或LP结构写成推断值。对于名称中含有基金、投资、创投、资本、合伙等关键词且类型属于VC/PE或广义PE/VC的主体，统一进入补充核验队列。后续只有在PDF、基金业协会或工商资料中找到来源，才填入备案编码和GP/LP结构。",
    )

    add_heading(doc, "七、PostgreSQL结果与数据披露", 1)
    if not pg_summary.empty:
        add_para(
            doc,
            "为补充数据库结果，本周使用本机PostgreSQL 18工具链启动临时实例，将第十周生成的10张CSV结果表导入数据库后，再通过SQL统一导出查询结果。这样既披露了数据库运行结果，又避免在没有5432主库口令时伪造入库。",
        )
        add_table(
            doc,
            ["披露项目", "披露值", "来源说明"],
            pg_summary[["item", "value", "source_note"]].values.tolist(),
            [3.0, 3.2, 8.8],
        )
        add_table(
            doc,
            ["表名", "行数"],
            pg_counts[["table_name", "row_count"]].values.tolist(),
            [8.5, 2.0],
        )
        add_table(
            doc,
            ["代码", "公司", "板块", "PE/VC占比", "强度指数", "第一大股东", "有效股东数", "质量"],
            pg_company[
                [
                    "stock_code",
                    "company_short",
                    "market",
                    "broad_pevc_record_share",
                    "pevc_strength_index",
                    "top1_ratio_pct",
                    "effective_shareholder_count",
                    "quality_gate",
                ]
            ].values.tolist(),
            [1.5, 1.8, 1.3, 1.8, 1.8, 1.8, 1.8, 1.5],
        )
        add_table(
            doc,
            ["主体类型", "画像行数", "来源记录数", "公司出现次数", "广义PE/VC主体", "基金/平台主体"],
            pg_investor_type[
                [
                    "investor_type_final",
                    "investor_profile_rows",
                    "source_record_count",
                    "company_mentions",
                    "broad_pevc_profiles",
                    "fund_like_profiles",
                ]
            ].values.tolist(),
            [2.8, 1.7, 1.9, 2.0, 2.1, 2.1],
        )
        add_table(
            doc,
            ["优先级", "队列行数", "备案编码已填", "GP已填", "LP已填", "披露状态"],
            pg_fund_status[
                [
                    "manual_priority",
                    "fund_queue_rows",
                    "amac_code_filled",
                    "gp_filled",
                    "lp_structure_filled",
                    "disclosure_status",
                ]
            ].values.tolist(),
            [3.5, 1.6, 1.9, 1.5, 1.5, 4.8],
        )
        add_para(
            doc,
            "上表显示，基金队列中备案编码、GP和LP结构仍为0条已填，这不是漏填，而是遵循“PDF/三表未披露就留空”的工程化原则。后续可在保留该空值边界的基础上，使用基金业协会、工商资料或PDF原文补证。",
        )
    else:
        add_para(
            doc,
            "当前尚未运行临时PostgreSQL导入脚本，因此本节暂只保留数据库导入脚本和主库连接审计。运行database/run_postgres_temp_week10.ps1后会自动生成披露表。",
        )

    add_heading(doc, "八、验证结果与不足", 1)
    validation_display = validation[["check_item", "status", "result", "note"]].copy()
    validation_display["status"] = validation_display["status"].map(display_status)
    val_rows = validation_display.values.tolist()
    add_table(doc, ["检查项", "状态", "结果", "说明"], val_rows, [3.0, 2.0, 4.5, 5.5])
    add_para(
        doc,
        "本周不足主要有三点：第一，PostgreSQL主库5432仍受本机口令限制，本周采用临时实例完成真实导入和查询披露，后续仍需迁移到本人长期数据库；第二，样本只有8家公司，描述统计和OLS只能作为探索性展示；第三，PE基金深度字段仍需外部核验，当前报告只完成队列化管理，没有完成主体穿透。",
    )

    add_heading(doc, "九、下一步计划", 1)
    next_rows = [
        ["数据库", "把临时PostgreSQL导入流程迁移到本人长期库，保存表行数校验日志", "真实入库记录"],
        ["扩样", "继续增加同口径公司，优先补足各板块样本数", "更稳定的描述统计"],
        ["主体穿透", "对P1基金队列补充备案编码、GP、LP结构和来源", "基金深度字段"],
        ["研究设计", "将PE/VC强度、股权分散度和控制变量整理为可回归面板", "正式研究问题草案"],
    ]
    add_table(doc, ["方向", "具体动作", "预期产出"], next_rows, [2.5, 8.0, 4.5])

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
