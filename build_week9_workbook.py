"""Build the Week 9 Excel workbook with a plain Python dependency stack.

The JavaScript workbook builder is kept for Codex visual preview only. This
script is the default reproducible entry point for GitHub/teacher-side checks.
"""

from __future__ import annotations

import json
from copy import copy
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUTPUTS = ROOT / "outputs"
PAYLOAD_PATH = OUTPUTS / "week9_workbook_data.json"
OUTPUT_PATH = OUTPUTS / "week9_summary_workbook.xlsx"


LABELS = {
    "queue_id": "编号",
    "stock_code": "证券代码",
    "company_short": "公司简称",
    "market": "板块",
    "source": "来源",
    "issue_type": "问题类型",
    "week8_status": "第八周状态",
    "week9_resolution": "第九周处理结论",
    "data_action": "数据处理动作",
    "analysis_action": "分析处理动作",
    "evidence_principle": "证据原则",
    "remaining_risk": "剩余风险",
    "has_vc_or_pe": "是否有VC/PE",
    "vc_record_count": "VC记录数",
    "pe_record_count": "PE记录数",
    "broad_pevc_record_count": "广义PE/VC记录数",
    "broad_pevc_record_share": "广义PE/VC记录占比",
    "pevc_intensity_level": "PE/VC强度",
    "distinct_investor_type_count": "投资主体类型数",
    "week8_p1_item_count": "第八周P1项",
    "week9_closed_p1_count": "第九周闭环P1项",
    "week9_unresolved_p1_count": "第九周未闭环P1项",
    "transfer_event_count": "转让事件数",
    "top1_ratio_pct": "第一大股东比例",
    "top3_ratio_pct": "前三大股东比例",
    "hhi": "HHI",
    "effective_shareholder_count": "有效股东数",
    "broad_pevc_holder_count": "广义PE/VC股东数",
    "broad_pevc_ratio_sum_pct": "广义PE/VC持股比例",
    "dispersion_level": "分散度判断",
    "week9_research_note": "研究使用说明",
    "selected_time_point": "选取时点",
    "shareholder_count": "股东数",
    "analysis_note": "分析说明",
    "company_count": "公司数",
    "vc_pe_supported_count": "有VC/PE公司数",
    "avg_broad_pevc_record_share": "平均PE/VC记录占比",
    "avg_distinct_investor_type_count": "平均主体类型数",
    "avg_top1_ratio_pct": "平均第一大股东比例",
    "avg_effective_shareholder_count": "平均有效股东数",
    "week9_interpretation": "口径说明",
    "investor_type": "投资主体类型",
    "table_role": "表内角色",
    "record_count": "记录数",
    "is_broad_pevc": "是否广义PE/VC",
    "research_question": "研究问题",
    "current_conclusion": "当前观察",
    "check_name": "检查项",
    "status": "状态",
    "detail": "说明",
}


HEADER_FILL = PatternFill("solid", fgColor="1F4D78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
HEADER_FONT = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
SUBHEADER_FONT = Font(name="Microsoft YaHei", bold=True, color="17365D")
BODY_FONT = Font(name="Microsoft YaHei", size=10)
THIN_GRAY = Side(style="thin", color="CBD5E1")


def text_stock_code(value: object) -> str:
    return str(value or "").strip().zfill(6)


def normalize(value: object, header: str = "") -> object:
    if header == "stock_code":
        return text_stock_code(value)
    if value == "":
        return None
    return value


def set_common_style(ws) -> None:
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows():
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def style_table(ws, min_row: int, max_row: int, min_col: int, max_col: int, dark_header: bool = True) -> None:
    fill = HEADER_FILL if dark_header else SUBHEADER_FILL
    font = HEADER_FONT if dark_header else SUBHEADER_FONT
    for cell in ws[min_row][min_col - 1 : max_col]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    border = Border(top=THIN_GRAY, bottom=THIN_GRAY, left=THIN_GRAY, right=THIN_GRAY)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = border


def write_table(ws, records: list[dict], headers: list[str], widths: dict[str, float] | None = None) -> None:
    widths = widths or {}
    ws.append([LABELS.get(header, header) for header in headers])
    for row in records:
        ws.append([normalize(row.get(header), header) for header in headers])
    set_common_style(ws)
    if ws.max_row >= 1:
        style_table(ws, 1, ws.max_row, 1, len(headers), dark_header=True)
    for idx, header in enumerate(headers, 1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = widths.get(header, 16)
        if header == "stock_code":
            for cell in ws[letter]:
                cell.number_format = "@"


def add_summary_sheet(wb: Workbook, payload: dict) -> None:
    ws = wb.active
    ws.title = "总览"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:H1")
    ws["A1"] = "霍泓锟第九周任务总览"
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=16)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:H2")
    ws["A2"] = "P1复核闭环、PostgreSQL导入准备与PE/VC研究变量形成"
    ws["A2"].font = Font(name="Microsoft YaHei", color="4B5563")
    ws.row_dimensions[2].height = 22

    ws.append([])
    ws.append(["指标", "数值", "说明"])
    summary_rows = [
        ["样本公司", payload["summary"]["company_count"], "沿用统一8家公司样本"],
        ["认缴/增资记录", payload["summary"]["subscription_records"], "来自第八周清洗表"],
        ["股权快照记录", payload["summary"]["snapshot_records"], "来自第八周清洗表"],
        ["股权转让记录", payload["summary"]["transfer_records"], "来自第八周清洗表"],
        ["第八周P1复核项", payload["summary"]["week8_p1_items"], "本周逐条闭环处理"],
        ["第九周已闭环P1项", payload["summary"]["week9_closed_p1"], "闭环不等于补数"],
        ["第九周未闭环P1项", payload["summary"]["week9_unresolved_p1"], "当前为0"],
        ["研究变量公司数", payload["summary"]["research_company_count"], "用于描述性研究问题"],
        ["PostgreSQL状态", payload["summary"]["postgres_status"], "服务可用，导入需密码"],
    ]
    for row in summary_rows:
        ws.append(row)

    start = 4
    end = start + len(summary_rows)
    style_table(ws, start, end, 1, 3, dark_header=False)

    board_start_col = 5
    board_headers = ["板块", "公司数", "平均PE/VC占比", "平均有效股东数"]
    for offset, value in enumerate(board_headers):
        ws.cell(row=4, column=board_start_col + offset, value=value)
    for r, row in enumerate(payload["board_stats"], 5):
        ws.cell(r, 5, row["market"])
        ws.cell(r, 6, row["company_count"])
        ws.cell(r, 7, row["avg_broad_pevc_record_share"])
        ws.cell(r, 8, row["avg_effective_shareholder_count"])
    style_table(ws, 4, 4 + len(payload["board_stats"]), 5, 8, dark_header=False)

    for col, width in {"A": 24, "B": 24, "C": 54, "E": 18, "F": 18, "G": 18, "H": 18}.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows():
        for cell in row:
            font = copy(cell.font)
            font.name = "Microsoft YaHei"
            cell.font = font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A4"


def add_table_sheet(
    wb: Workbook,
    name: str,
    records: list[dict],
    headers: list[str],
    widths: dict[str, float] | None = None,
) -> None:
    ws = wb.create_sheet(name)
    write_table(ws, records, headers, widths)


def main() -> None:
    payload = json.loads(PAYLOAD_PATH.read_text(encoding="utf-8"))
    wb = Workbook()
    add_summary_sheet(wb, payload)

    add_table_sheet(
        wb,
        "复核闭环",
        payload["review_resolved"],
        [
            "queue_id",
            "stock_code",
            "company_short",
            "market",
            "issue_type",
            "week8_status",
            "week9_resolution",
            "data_action",
            "analysis_action",
            "evidence_principle",
            "remaining_risk",
        ],
        {
            "queue_id": 16,
            "company_short": 16,
            "issue_type": 24,
            "week9_resolution": 24,
            "data_action": 38,
            "analysis_action": 44,
            "evidence_principle": 24,
            "remaining_risk": 44,
        },
    )

    add_table_sheet(
        wb,
        "研究变量",
        payload["research_variables"],
        [
            "stock_code",
            "company_short",
            "market",
            "has_vc_or_pe",
            "broad_pevc_record_share",
            "pevc_intensity_level",
            "distinct_investor_type_count",
            "week8_p1_item_count",
            "week9_closed_p1_count",
            "week9_unresolved_p1_count",
            "top1_ratio_pct",
            "hhi",
            "effective_shareholder_count",
            "dispersion_level",
            "week9_research_note",
        ],
        {
            "stock_code": 14,
            "company_short": 16,
            "market": 12,
            "week9_research_note": 42,
        },
    )

    add_table_sheet(
        wb,
        "股权分散度",
        payload["ownership_metrics"],
        [
            "stock_code",
            "company_short",
            "market",
            "selected_time_point",
            "shareholder_count",
            "top1_ratio_pct",
            "top3_ratio_pct",
            "hhi",
            "effective_shareholder_count",
            "broad_pevc_holder_count",
            "broad_pevc_ratio_sum_pct",
            "dispersion_level",
            "analysis_note",
        ],
        {
            "stock_code": 14,
            "company_short": 16,
            "selected_time_point": 24,
            "analysis_note": 38,
        },
    )

    add_table_sheet(
        wb,
        "板块统计",
        payload["board_stats"],
        [
            "market",
            "company_count",
            "vc_pe_supported_count",
            "avg_broad_pevc_record_share",
            "avg_distinct_investor_type_count",
            "avg_top1_ratio_pct",
            "avg_effective_shareholder_count",
            "week9_interpretation",
        ],
        {"week9_interpretation": 48},
    )

    add_table_sheet(
        wb,
        "投资主体矩阵",
        payload["investor_matrix"],
        ["market", "investor_type", "table_role", "record_count", "is_broad_pevc"],
        {"investor_type": 24, "table_role": 16},
    )

    add_table_sheet(
        wb,
        "研究问题观察",
        payload["research_rows"],
        [
            "research_question",
            "stock_code",
            "company_short",
            "market",
            "pevc_intensity_level",
            "broad_pevc_record_share",
            "top1_ratio_pct",
            "effective_shareholder_count",
            "dispersion_level",
            "current_conclusion",
        ],
        {"research_question": 48, "stock_code": 14, "company_short": 16, "current_conclusion": 44},
    )

    add_table_sheet(
        wb,
        "数据库状态",
        payload["postgres_audit"],
        ["check_name", "status", "detail"],
        {"check_name": 20, "status": 24, "detail": 72},
    )

    OUTPUTS.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
