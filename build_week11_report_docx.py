"""Build the Week 11 DOCX report."""

from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DERIVED = ROOT / "data" / "derived"
DATABASE = ROOT / "database"
REPORT = ROOT / "report"
OUTPUT = REPORT / "霍泓锟_第十一周任务报告.docx"
RUN_DATE = "2026年9月2日"


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def set_east_asian_font(run, name: str = "Microsoft YaHei") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_para(doc: Document, text: str, size: float = 10.5, bold: bool = False, align=None) -> None:
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_east_asian_font(run)
    run.font.size = Pt(size)
    run.font.bold = bold


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_east_asian_font(run)
    run.font.bold = True
    run.font.size = Pt(14 if level == 1 else 12)
    run.font.color.rgb = RGBColor(31, 77, 120)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: object, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("" if text is None else str(text))
    set_east_asian_font(run)
    run.font.size = Pt(8.5)
    run.font.bold = bold


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D0D7DE")


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, header, bold=True)
        set_cell_shading(cell, "E8EEF5")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = Cm(widths[idx])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cells[idx].width = Cm(widths[idx])
    set_table_borders(table)


def get_metric(quality: list[dict[str, str]], metric: str) -> str:
    for row in quality:
        if row.get("metric") == metric:
            return row.get("value", "")
    return ""


def main() -> None:
    REPORT.mkdir(parents=True, exist_ok=True)
    quality = read_csv(DERIVED / "week11_quality_dashboard.csv")
    db_plan = read_csv(DERIVED / "week11_database_migration_plan.csv")
    readiness = read_csv(DERIVED / "week11_persistent_db_readiness.csv")
    fund_queue = read_csv(DERIVED / "week11_fund_deep_enrichment_queue.csv")
    manual = read_csv(DERIVED / "week11_manual_verification_template.csv")
    research = read_csv(DERIVED / "week11_research_design_matrix.csv")
    expansion = read_csv(DERIVED / "week11_sample_expansion_plan.csv")
    pg_disclosure = read_csv(DATABASE / "postgresql_week11_disclosure.csv")
    completion = read_csv(DERIVED / "week11_completion_checklist.csv")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("霍泓锟 第十一周任务报告")
    set_east_asian_font(run)
    run.font.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(23, 54, 93)

    add_para(doc, f"主题：数据库长期化迁移准备与基金深度字段核验 | 提交日期：{RUN_DATE}", 10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "本周任务并不是重新追求文件数量，而是在第十周研究面板和PostgreSQL临时库结果基础上，把后续最容易被老师追问的两个环节做实：一是长期库如何接入，二是基金备案编码、GP和LP结构如何在不伪造的前提下继续补全。", 10.5)

    add_heading(doc, "一、本周承接基础与核心结果")
    add_para(doc, "第十周已经形成统一8家公司研究面板、192条投资主体画像和32条基金核验队列；临时PostgreSQL实例已验证10张表共406行可以导入。本周继续沿用这些真实输出，不从人工Final反推Auto，也不把未披露字段改写为推测值。")
    add_table(
        doc,
        ["指标", "数值", "解释"],
        [[row.get("metric", ""), row.get("value", ""), row.get("interpretation", "")] for row in quality],
        [4.0, 2.5, 10.5],
    )

    add_heading(doc, "二、PostgreSQL结果与长期库迁移")
    add_para(doc, "第十一周将数据库工作拆成两层：临时库用于验证代码和CSV能否复现；长期库用于后续连续周次沉淀数据。当前本机5432端口可检测到服务，但无交互运行需要PGPASSWORD，因此长期库导入脚本已经准备好，是否真正写入主库要以口令和权限为前提。")
    add_table(
        doc,
        ["任务", "状态", "证据", "下一步"],
        [[row.get("task", ""), row.get("status", ""), row.get("evidence", ""), row.get("next_action", "")] for row in db_plan],
        [3.3, 3.0, 6.2, 4.5],
    )
    if pg_disclosure:
        add_para(doc, "临时PostgreSQL复现后导出的本周披露结果如下：")
        add_table(
            doc,
            ["项目", "值", "说明"],
            [[row.get("item", ""), row.get("value", ""), row.get("source_note", "")] for row in pg_disclosure],
            [4.2, 3.0, 9.8],
        )
    else:
        add_para(doc, "当前尚未检测到第十一周临时库披露结果，可运行 database/run_postgres_temp_week11.ps1 生成。")

    add_heading(doc, "三、基金深度字段核验设计")
    add_para(doc, "老师强调PE基金深度提取时，不能只停留在投资者名称和持股比例，还要继续追备案编码、GP和LP结构。但这些字段在招股书三表中经常不直接披露，所以本周采用“队列化核验”：先锁定主体、公司、页码和优先级，再由人工或外部来源逐条确认。")
    add_table(
        doc,
        ["记录ID", "公司", "投资主体", "类型", "页码", "需核字段", "处理原则"],
        [
            [
                row.get("record_id", ""),
                row.get("company_short", ""),
                row.get("investor_name", ""),
                row.get("investor_type_final", ""),
                row.get("pdf_pages_observed", ""),
                "AMAC/GP/LP",
                row.get("engineering_rule", ""),
            ]
            for row in fund_queue[:12]
        ],
        [1.4, 1.8, 5.6, 1.6, 2.0, 2.0, 4.2],
    )
    add_para(doc, f"本周基金深度队列共 {len(fund_queue)} 条，人工核验事项共 {len(manual)} 条。当前备案编码、GP和LP结构均未由PDF三表直接披露，因此保持空值，并在模板中标记为待核验。")

    add_heading(doc, "四、研究问题与未来扩展")
    add_para(doc, "第十一周的研究设计不急于扩大结论，而是把变量、样本、方法和风险先固定下来。等长期库导入成功、基金深度字段补充完成后，再推进正式的描述统计、相关分析和稳健性检验。")
    add_table(
        doc,
        ["问题", "核心变量", "方法", "当前限制"],
        [
            [
                row.get("research_question", ""),
                row.get("key_explanatory_variable", ""),
                row.get("method", ""),
                row.get("current_limit", ""),
            ]
            for row in research
        ],
        [5.2, 4.2, 3.2, 5.2],
    )
    add_table(
        doc,
        ["计划", "范围", "目标", "验收标准", "状态"],
        [[row.get("period", ""), row.get("scope", ""), row.get("target", ""), row.get("success_criteria", ""), row.get("status", "")] for row in expansion],
        [2.0, 3.0, 6.0, 5.0, 2.0],
    )

    add_heading(doc, "五、提交文件与可复现说明")
    add_para(doc, "本周提交包保留统一入口 code/run_all_week11_codex.ps1。运行顺序为：生成派生表、启动临时PostgreSQL并导入、重新写入PG披露结果、生成Excel、生成Word和PDF报告。长期库脚本单独放在 database/run_postgres_persistent_week11.ps1，避免把无密码状态误报为导入成功。")
    add_table(
        doc,
        ["项目", "状态", "证据"],
        [[row.get("item", ""), row.get("status", ""), row.get("evidence", "")] for row in completion],
        [5.0, 2.2, 9.8],
    )

    add_heading(doc, "六、不足与第十二周计划")
    add_para(doc, "第一，当前长期PostgreSQL主库仍需要口令后才能实际导入，临时库只能证明脚本和表结构可复现。第二，基金备案编码、GP和LP结构尚未完成外部核验，本周坚持留空原则，因此深度字段完整率暂时较低。第三，目前8家公司仍偏小，探索性OLS和相关系数只能作为流程演示，不能夸大为稳健结论。")
    add_para(doc, "第十二周应优先完成三项工作：设置PGPASSWORD并把第十一周派生表写入长期库；按P1优先级核验基金备案编码、GP和LP结构；在四个板块继续补样，并确保新增样本沿用PDF/Markdown、Auto、Gold/Final和Cross-check的分层流程。")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("霍泓锟 第十一周任务提交")
    set_east_asian_font(run)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(90, 90, 90)

    doc.save(OUTPUT)
    print(f"saved {OUTPUT}")


if __name__ == "__main__":
    main()
