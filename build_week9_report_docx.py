"""Build Week 9 DOCX report from pipeline outputs."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUTS = ROOT / "outputs"
REPORT = ROOT / "report"

FONT_CN = "宋体"
FONT_HEI = "黑体"
FONT_EN = "Times New Roman"
BLUE = RGBColor(31, 78, 121)


def set_run_font(run, cn_font=FONT_CN, size=10.5, bold=None, color=None):
    run.font.name = FONT_EN
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn_font)
    rfonts.set(qn("w:ascii"), FONT_EN)
    rfonts.set(qn("w:hAnsi"), FONT_EN)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_border(cell, edge, val="nil", sz="0", color="auto"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    element = borders.find(qn(f"w:{edge}"))
    if element is None:
        element = OxmlElement(f"w:{edge}")
        borders.append(element)
    element.set(qn("w:val"), val)
    element.set(qn("w:sz"), sz)
    element.set(qn("w:space"), "0")
    element.set(qn("w:color"), color)


def apply_three_line_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    rows = len(table.rows)
    if rows == 0:
        return
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
                set_cell_border(cell, edge, "nil", "0")
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                for run in para.runs:
                    set_run_font(run, size=8.2)
    for cell in table.rows[0].cells:
        set_cell_border(cell, "top", "single", "12", "000000")
        set_cell_border(cell, "bottom", "single", "8", "000000")
        for para in cell.paragraphs:
            for run in para.runs:
                set_run_font(run, cn_font=FONT_HEI, size=8.5, bold=True)
    for cell in table.rows[-1].cells:
        set_cell_border(cell, "bottom", "single", "12", "000000")


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, cn_font=FONT_HEI, size=18, bold=True, color=BLUE)
    p.paragraph_format.space_after = Pt(8)


def add_heading(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, cn_font=FONT_HEI, size=13.5, bold=True, color=BLUE)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(5)
    for part in str(text).split("\n"):
        if p.runs:
            p.add_run().add_break()
        r = p.add_run(part)
        set_run_font(r, size=10.5)
    return p


def add_small_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = "" if val is None else str(val)
    apply_three_line_table(table)
    return table


def as_rows(records, fields, limit=None):
    out = []
    for row in records[:limit]:
        out.append([row.get(f, "") for f in fields])
    return out


def main():
    payload = json.loads((OUTPUTS / "week9_workbook_data.json").read_text(encoding="utf-8"))
    summary = payload["summary"]
    board_stats = payload["board_stats"]
    research_vars = payload["research_variables"]
    review_resolved = payload["review_resolved"]
    postgres_audit = payload["postgres_audit"]

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = FONT_EN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(10.5)

    add_title(doc, "第九周任务报告：P1复核闭环、数据库准备与初步研究问题形成")
    add_para(doc, "姓名：霍泓锟\n日期：2026年8月19日\n主题：PE/VC招股书三表数据的复核闭环、数据库化和研究变量准备")

    add_heading(doc, "一、本周任务定位")
    add_para(
        doc,
        "第九周承接此前未来一周与未来一个月任务规划，也延续第八周报告中提出的遗留事项。本周重点不再是继续增加抽取文件数量，而是把已经形成的8家公司三表数据继续向后推进：一是关闭P1人工复核队列；二是补齐PostgreSQL导入所需的表结构、导入脚本和连接审计；三是在第八周研究变量草案基础上提出一个可以继续扩展的小型研究问题。",
    )
    add_para(
        doc,
        f"主流程读取认缴/增资记录{summary['subscription_records']}条、股权快照记录{summary['snapshot_records']}条、股权转让记录{summary['transfer_records']}条，继续覆盖统一8家公司样本。本周所有派生结果都从第八周清洗表重新生成，未从Gold或Final之外手工拼接统计值。",
    )

    add_heading(doc, "二、P1复核队列处理")
    add_para(
        doc,
        f"第八周留下{summary['week8_p1_items']}个P1复核项，第九周已闭环{summary['week9_closed_p1']}个，未闭环{summary['week9_unresolved_p1']}个。这里的“闭环”不是把缺失值补齐，而是把每条问题的处理原则、数据动作和分析动作写清楚。黄山谷捷若PDF没有披露完整可求和比例，继续留空；赛分科技670.6317%的比例合计异常被识别为同一根因的重复命中，在派生统计中剔除异常时点。",
    )
    add_small_table(
        doc,
        ["编号", "公司", "问题类型", "处理结论", "数据动作"],
        as_rows(review_resolved, ["queue_id", "company_short", "issue_type", "week9_resolution", "data_action"], limit=10),
    )

    add_heading(doc, "三、数据库化推进")
    add_para(
        doc,
        "本周生成了PostgreSQL表结构、导入脚本和运行模板。数据库设计不直接把三张大表揉成一张宽表，而是保留公司维度表、复核闭环表、研究变量表、板块统计表和股权分散度表，便于之后继续追加公司样本或把基金备案编码、GP/LP结构单独拆表。",
    )
    add_small_table(doc, ["检查项", "状态", "说明"], as_rows(postgres_audit, ["check_name", "status", "detail"]))
    add_para(
        doc,
        "本机PostgreSQL服务可用，localhost:5432处于接受连接状态；但当前没有提供免密口令，因此本周不伪造“已经写入数据库”的结果，而是记录为READY_NEED_PASSWORD。后续只需设置PGPASSWORD并运行database/run_postgres_import_week9.ps1，即可执行导入并得到各表行数。",
    )

    add_heading(doc, "四、初步研究问题")
    add_para(
        doc,
        "本周提出的小型研究问题是：不同板块PE/VC进入强度与上市前股权分散度是否存在差异？变量来源包括investor_type_final分类、广义PE/VC记录占比、最新可求和股权快照中的第一大股东比例、HHI和有效股东数。由于当前只有8家公司，这一部分只做描述性分析，不做显著性检验或正式回归。",
    )
    add_small_table(
        doc,
        ["代码", "公司", "板块", "PE/VC强度", "PE/VC记录占比", "第一大股东比例", "有效股东数", "说明"],
        as_rows(
            research_vars,
            [
                "stock_code",
                "company_short",
                "market",
                "pevc_intensity_level",
                "broad_pevc_record_share",
                "top1_ratio_pct",
                "effective_shareholder_count",
                "week9_research_note",
            ],
        ),
    )

    add_heading(doc, "五、板块描述性统计")
    add_para(
        doc,
        "从板块看，当前样本在主板、创业板、科创板和北交所各有2家公司，适合做非常初步的横向比较。这里的PE/VC进入强度是记录占比口径，会受到招股书披露详略影响；股权分散度则只使用通过比例合计校验的股权快照。",
    )
    add_small_table(
        doc,
        ["板块", "公司数", "有VC/PE记录公司数", "平均PE/VC记录占比", "平均第一大股东比例", "平均有效股东数"],
        as_rows(
            board_stats,
            [
                "market",
                "company_count",
                "vc_pe_supported_count",
                "avg_broad_pevc_record_share",
                "avg_top1_ratio_pct",
                "avg_effective_shareholder_count",
            ],
        ),
    )

    add_heading(doc, "六、本周不足与下一步")
    add_para(
        doc,
        "本周仍有三个不足。第一，PostgreSQL服务虽然可用，但因为没有提供本机数据库密码，导入停留在可执行脚本和连接审计层面；这需要下一周补充真实导入日志。第二，样本只有8家公司，描述性统计只能帮助提出研究问题，不能支撑稳健结论。第三，复杂PE基金的备案编码、GP/LP结构和政府基金属性仍未系统补齐，后续需要把基金主体继续拆成更细的数据库表。",
    )
    add_para(
        doc,
        "下一步建议优先做三件事：其一，补充PostgreSQL密码并执行真实导入，保存行数校验结果；其二，扩大样本后重新计算PE/VC参与强度与股权分散度；其三，对PE、VC、政府基金、产业资本和员工持股平台分别配置分类规则，使研究变量从“粗分类可用”推进到“细分类可解释”。",
    )

    out_path = REPORT / "霍泓锟_第九周任务报告.docx"
    doc.save(out_path)
    print(out_path)


if __name__ == "__main__":
    main()
