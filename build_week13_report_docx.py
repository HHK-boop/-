from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DERIVED = ROOT / "data" / "derived"
OUTPUTS = ROOT / "outputs"
REPORT_DIR = ROOT / "report"
REPORT_DIR.mkdir(parents=True, exist_ok=True)

REPORT_PATH = REPORT_DIR / "霍泓锟_第十三周扩样抽取与人工核验报告.docx"
CHART_PATH = OUTPUTS / "week13_event_candidates.png"

NAVY = "1F4E78"
BLUE = "D9EAF7"
PALE = "F4F7FA"
GRAY = "D9E2F3"
TEXT = RGBColor(31, 41, 55)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def load_data():
    summary = json.loads((OUTPUTS / "week13_summary.json").read_text(encoding="utf-8"))
    return {
        "summary": summary,
        "companies": read_csv(DERIVED / "week13_company_progress.csv"),
        "events": read_csv(DERIVED / "week13_event_summary.csv"),
        "crosscheck": read_csv(DERIVED / "week13_old_new_crosscheck.csv"),
        "review": read_csv(DERIVED / "week13_manual_review_queue.csv"),
        "validation": read_csv(ROOT / "validation" / "week13_validation_summary.csv"),
        "pg": read_csv(ROOT / "database" / "postgresql_week13_disclosure.csv"),
        "sources": read_csv(DERIVED / "week13_selected_companies.csv"),
        "profiles": read_csv(DERIVED / "week13_investor_profile.csv"),
    }


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key in ["val", "sz", "space", "color"]:
            if key in kwargs[edge]:
                element.set(qn("w:" + key), str(kwargs[edge][key]))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_table_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_borders_and_font(table, header=True, font_size=8.5):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        keep_table_row(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(
                cell,
                top={"val": "single", "sz": 4, "color": GRAY},
                bottom={"val": "single", "sz": 4, "color": GRAY},
                left={"val": "single", "sz": 4, "color": GRAY},
                right={"val": "single", "sz": 4, "color": GRAY},
            )
            if header and row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                    run.font.size = Pt(font_size)
                    if header and row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
        if header and row_index == 0:
            set_repeat_table_header(row)


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = str(header)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = "" if value is None else str(value)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    add_borders_and_font(table, font_size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_note(doc, text):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = text
    set_cell_shading(cell, BLUE)
    set_cell_border(cell, left={"val": "single", "sz": 16, "color": NAVY})
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.size = Pt(9.5)
            run.font.color.rgb = TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.add_run(text)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.add_run(text)
    return paragraph


def add_manual_step(doc, number, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.55)
    paragraph.paragraph_format.first_line_indent = Cm(-0.45)
    paragraph.add_run(f"{number}. ").bold = True
    paragraph.add_run(text)
    return paragraph


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def build_chart(events):
    labels = [row["event_type"] for row in events]
    counts = [int(row["candidate_records"]) for row in events]
    high = [int(row["high_confidence_records"]) for row in events]
    width, height = 1550, 650
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    bold_path = Path(r"C:\Windows\Fonts\msyhbd.ttc")
    font = ImageFont.truetype(str(font_path), 28)
    small = ImageFont.truetype(str(font_path), 23)
    bold = ImageFont.truetype(str(bold_path if bold_path.exists() else font_path), 34)
    draw.text((55, 28), "第十三周事件候选分布", font=bold, fill="#1F2937")
    draw.rectangle((985, 36, 1020, 61), fill="#5B9BD5")
    draw.text((1032, 31), "保留Auto候选", font=small, fill="#1F2937")
    draw.rectangle((1242, 36, 1277, 61), fill="#1F4E78")
    draw.text((1289, 31), "高置信候选", font=small, fill="#1F2937")
    left, top, bar_width = 250, 115, 1180
    max_count = max(counts) or 1
    row_height = 92
    for index, (label, count, high_count) in enumerate(zip(labels, counts, high)):
        y = top + index * row_height
        draw.text((55, y + 12), label, font=font, fill="#1F2937")
        draw.line((left, y + 66, left + bar_width, y + 66), fill="#E5E7EB", width=2)
        count_width = int(bar_width * count / max_count)
        high_width = int(bar_width * high_count / max_count)
        draw.rectangle((left, y + 8, left + count_width, y + 43), fill="#5B9BD5")
        draw.rectangle((left, y + 45, left + high_width, y + 64), fill="#1F4E78")
        draw.text((left + count_width + 10, y + 7), str(count), font=small, fill="#1F2937")
        draw.text((left + high_width + 10, y + 42), str(high_count), font=small, fill="#1F4E78")
    image.save(CHART_PATH, format="PNG")


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color in [("Title", 24, NAVY), ("Heading 1", 16, NAVY), ("Heading 2", 13, NAVY), ("Heading 3", 11, "000000")]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "PEVC招股说明书结构化解析项目｜第十三周"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(102, 112, 122)
        add_page_number(section.footer.paragraphs[0])


def add_cover(doc):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(64)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("PEVC招股说明书项目")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(22)
    run = paragraph.add_run("第十三周扩样抽取与人工核验报告")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(25)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(10)
    run = paragraph.add_run("从56家公司处理队列到首批12家北交所样本的可复现落地")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(75, 85, 99)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = [("姓名", "霍泓锟"), ("重点市场", "北京证券交易所"), ("报告周期", "第十三周"), ("提交日期", "2026年9月16日")]
    for row, values_row in zip(table.rows, values):
        row.cells[0].text, row.cells[1].text = values_row
        row.cells[0].width = Cm(4)
        row.cells[1].width = Cm(8)
        set_cell_shading(row.cells[0], BLUE)
        for cell in row.cells:
            set_cell_border(cell, bottom={"val": "single", "sz": 6, "color": GRAY})
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                    run.font.size = Pt(11)
    doc.add_paragraph().paragraph_format.space_before = Pt(115)
    paragraph = doc.add_paragraph("本报告中的Auto记录均为机器定位或字段候选；只有经过PDF原文核验后才能进入Gold/Final。")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(107, 114, 128)
    doc.add_page_break()


def build_report():
    data = load_data()
    summary = data["summary"]
    build_chart(data["events"])

    doc = Document()
    configure_document(doc)
    add_cover(doc)

    add_heading(doc, "摘要", level=1)
    doc.add_paragraph(
        "第十三周承接第十二周形成的298家公司来源目录和56家公司处理队列，工作重点从“规划扩样”推进到“首批扩样实际运行”。"
        "本周从北交所P1队列中按关键词命中数和本地页码化文本可用性选取12家公司，独立读取招股书页码化文本，完成章节定位、事件候选抽取、投资主体类型候选、人工复核队列、旧新定位交叉校验和PostgreSQL入库验证。"
    )
    doc.add_paragraph(
        f"结果显示：12家公司共处理{summary['source_page_count']}页文本，形成{summary['auto_candidate_record_count']}条Auto候选；"
        f"规则排除{summary['excluded_transfer_boilerplate_count']}条“全国股转系统/公开转让”语境误命中后，保留{summary['retained_candidate_record_count']}条待核候选；"
        f"去重得到{summary['unique_investor_candidate_count']}个投资主体候选，并建立{summary['manual_review_queue_count']}条人工复核队列。"
        "旧定位页平均找回率为77.08%，该指标只反映定位结果的一致性，不代表Gold准确率。"
    )
    add_note(doc, "工程口径：Auto流程不读取Gold/Final的主体、页码或记录ID；招股书没有明确披露备案编码、GP或LP结构时保持空值。")

    add_heading(doc, "一、本周目标与完成情况", level=1)
    add_table(
        doc,
        ["目标", "执行结果", "状态"],
        [
            ["从扩样队列选择首批公司", "选择12家北交所公司，保留来源URL和页码化文本副本", "完成"],
            ["独立生成Auto候选", "从页码化文本定位5类事件，不读取Gold/Final", "完成"],
            ["完善investor_type分类", "输出7类规则和52个去重主体候选，保留人工确认状态", "完成"],
            ["建立自动与人工对比入口", "生成28条优先复核记录和空白人工结论字段", "完成"],
            ["数据库披露", "临时PostgreSQL导入11张表，共492行", "完成"],
            ["形成可提交成果", "Excel、Word、日志、SQL、CSV、源码和统一入口", "完成"],
        ],
        [4.5, 10.0, 2.0],
        9,
    )

    add_heading(doc, "二、样本选择与数据来源", level=1)
    doc.add_paragraph(
        "本周没有重新从互联网临时拼接样本，而是复用第十二周已建立的来源目录与处理队列。筛选条件依次为：板块为北交所、优先级为P1、"
        "本地存在带页码标记的招股书文本、关键词命中数较高。这样既延续前期计划，也保证统一入口可在离线状态重新运行。"
    )
    source_map = {row["stock_code"]: row for row in data["sources"]}
    company_rows = []
    for row in data["companies"]:
        source = source_map.get(row["stock_code"], {})
        company_rows.append([
            row["stock_code"], row["company_short"], row["page_count"],
            source.get("week12_keyword_hits", ""), row["chapter_groups_located"],
            row["retained_auto_candidates"], row["excluded_transfer_boilerplate"],
        ])
    doc.add_page_break()
    add_table(
        doc,
        ["代码", "公司", "页数", "W12命中", "章节组", "保留候选", "排除误命中"],
        company_rows,
        [2.0, 2.8, 1.4, 1.6, 1.6, 2.0, 2.2],
        8,
    )
    doc.add_paragraph(
        "每家公司在提交包内均保留来源URL、原始文本路径和包内文本路径。12家公司各保留180页页码化文本，共2160页；"
        "这批文本既是自动流程的输入，也是老师抽查数字来源时的直接证据。"
    )

    add_heading(doc, "三、可复现处理流程", level=1)
    add_number(doc, "输入隔离：主流程只读取第十二周处理队列和data/source_texts，不读取manual_gold或final目录。")
    add_number(doc, "页码解析：按“<!-- page:N -->”切分文本，建立公司、页码和正文之间的稳定索引。")
    add_number(doc, "章节定位：按历史沿革与股本形成、股东与股权结构、特殊投资条款、基金与备案四组规则定位重点页。")
    add_number(doc, "事件候选：分别抽取增资认缴、股权转让、股权快照、基金备案和特殊投资条款五类候选。")
    add_number(doc, "规则排除：对只出现“全国股转系统”“公开转让”且缺少转让方、受让方、回购、大宗交易等交易证据的记录做排除标记。")
    add_number(doc, "主体分类：输出investor_type候选、分类依据和人工状态；基金备案编码、GP、LP结构不根据名称猜测。")
    add_number(doc, "质量验证：生成旧新定位对比、28条人工复核队列、8项自动验收和PostgreSQL披露结果。")
    add_note(doc, "统一运行入口：在提交目录执行 PowerShell -ExecutionPolicy Bypass -File .\\code\\run_all_week13_codex.ps1。运行后可重新生成CSV、数据库披露、Excel和Word报告。")

    add_heading(doc, "四、Auto候选结果与解释", level=1)
    event_rows = [[row["event_type"], row["candidate_records"], row["companies_covered"], row["high_confidence_records"], row["with_date"], row["with_amount_or_shares"]] for row in data["events"]]
    add_table(
        doc,
        ["事件类型", "候选记录", "覆盖公司", "高置信", "含日期", "含金额/股数"],
        event_rows,
        [3.2, 2.0, 2.0, 2.0, 2.0, 2.6],
        8.5,
    )
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(CHART_PATH), width=Inches(6.35))
    caption = doc.add_paragraph("图1  第十三周事件候选分布（候选数不等于最终事件数）")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(107, 114, 128)

    doc.add_paragraph(
        "股权快照候选最多，为85条；股权转让候选70条；增资认缴63条；特殊投资条款38条；基金备案19条。"
        "这些数量用于安排人工复核工作量，不能直接作为公司真实融资轮次、投资事件数或准确率。"
        "同一页可能同时包含多个命中词，同一事件也可能跨页重复披露，因此Final阶段必须合并同义事件并以原文事实为准。"
    )

    doc.add_page_break()
    add_heading(doc, "五、投资主体类型与PE基金深度字段", level=1)
    doc.add_paragraph(
        "本周把investor_type作为重点字段，区分员工持股平台、政府基金/国资、VC、PE、CVC、企业法人候选和自然人。"
        "规则以名称线索与原文上下文共同判断；普通投资公司不能仅凭名称归为VC，企业法人股东也不能自动归为PE。"
    )
    type_counts = {}
    for row in data["profiles"]:
        key = row.get("investor_type_candidate", "未分类") or "未分类"
        type_counts[key] = type_counts.get(key, 0) + 1
    type_rows = [[key, value, "均为候选，需核对原文股东性质或基金属性"] for key, value in sorted(type_counts.items(), key=lambda item: (-item[1], item[0]))]
    add_table(doc, ["类型候选", "主体数", "使用边界"], type_rows, [4.0, 2.0, 10.0], 8.5)
    add_bullet(doc, "备案编码：只有招股书明确给出中国证券投资基金业协会备案编码时才填写。")
    add_bullet(doc, "GP：只有原文明确披露普通合伙人或执行事务合伙人时才填写。")
    add_bullet(doc, "LP结构：只有原文列示有限合伙人及其份额时才结构化，不根据工商常识补齐。")
    add_bullet(doc, "本周52个去重主体候选的上述深度字段均保持空值，体现“未披露就留空”，而不是数据遗漏。")

    add_heading(doc, "六、交叉校验与人工复核", level=1)
    doc.add_paragraph(
        "为了量化流程变化，本周将第十三周定位页与第三周旧候选页进行页码集合对比。12家公司旧定位页平均找回率为77.08%。"
        "该结果说明新规则能找回多数旧页，同时增加了一批新定位页；但页码重合不判断字段是否正确，因此不能称为准确率。"
    )
    cross_rows = [[row["stock_code"], row["company_short"], row["old_unique_pages"], row["week13_unique_pages"], row["overlap_pages"], f"{float(row['old_page_recovery_rate']):.1%}", row["newly_located_pages"]] for row in data["crosscheck"]]
    add_table(doc, ["代码", "公司", "旧页", "本周页", "重合页", "找回率", "新增页"], cross_rows, [1.8, 2.6, 1.6, 1.8, 1.8, 1.8, 1.8], 8)
    doc.add_paragraph(
        "人工复核队列共28条，超过保留候选的10%。抽样优先级不是纯随机：先纳入低/中置信、股权转让和基金备案候选，"
        "再补足各公司覆盖。人工操作时应在Excel“人工复核队列”中填写human_decision、corrected_value、reviewer和review_date。"
    )
    add_manual_step(doc, 1, "打开队列中的包内文本文件，跳转到source_page对应的页码标记。")
    add_manual_step(doc, 2, "确认是否为真实事件，尤其排除“申请挂牌并公开转让”等市场制度表述。")
    add_manual_step(doc, 3, "逐项核对日期、金额、股数、比例、转让方、受让方和事件关系，不从相邻事件拼接数字。")
    add_manual_step(doc, 4, "若招股书未披露字段，填空并在review_note记录“原文未披露”，不得按名称或常识推断。")
    add_manual_step(doc, 5, "人工确认后再把记录合并为Gold；Final保留修改前值、修改后值、原因和PDF页码。")

    add_heading(doc, "七、失败案例与当前不足", level=1)
    add_table(
        doc,
        ["问题", "本周表现", "处理方式", "仍需改进"],
        [
            ["转让词误命中", "26条记录仅描述全国股转系统或公开转让", "规则排除并保留排除原因", "继续加入交易主体和动词依存条件"],
            ["章节覆盖不完整", "迪尔化工只定位到3个章节组，基金备案组未命中", "不补造章节，进入人工检查", "允许目录别名和跨页表题识别"],
            ["主体名称噪声", "部分候选带括号、截断或OCR残片", "合并同名并过滤通用短语", "增加公司后缀词典和表格列约束"],
            ["同一事件重复", "同一事项可能在概览、历史沿革、股东表重复", "当前仍保留候选", "Gold阶段增加日期+主体+金额事件去重"],
            ["深度字段缺失", "备案编码、GP、LP在本批原文中未形成可靠披露", "全部留空并标注规则", "仅在原文证据充分时建派生表"],
            ["定位指标边界", "77.08%是旧页找回率，不是字段准确率", "报告和工作簿反复披露口径", "完成Gold复核后再计算Precision/Recall/F1"],
        ],
        [3.1, 4.2, 4.2, 4.2],
        8,
    )

    doc.add_page_break()
    add_heading(doc, "八、PostgreSQL验证与数据披露", level=1)
    doc.add_paragraph(
        "为验证结构化结果可以脱离Excel使用，本周在临时PostgreSQL实例中创建pevc_week13 schema，导入公司、章节定位、"
        "Auto候选、投资主体候选、人工队列、分类规则、交叉校验等11张表。数据库完成查询后已正常停止，不依赖常驻服务。"
    )
    pg_rows = [[row.get("item", ""), row.get("value", ""), row.get("source_note", "")] for row in data["pg"]]
    add_table(doc, ["披露项目", "结果", "来源/解释"], pg_rows, [4.3, 2.5, 9.0], 8.5)
    add_note(doc, "PostgreSQL共导入492行。备案编码、GP、LP已填写数均为0，数据库查询结果与“原文未披露就留空”的工程原则一致。")

    add_heading(doc, "九、质量验收", level=1)
    validation_rows = [[row["check_item"], row["status"], row["value"], row["rule"]] for row in data["validation"]]
    add_table(doc, ["检查项", "状态", "结果", "验收规则"], validation_rows, [4.0, 1.8, 3.0, 7.0], 8.2)
    doc.add_paragraph(
        f"8项自动检查全部通过。检查覆盖样本数、文本可读性、来源URL、章节定位、Gold隔离、误命中过滤、人工复核比例以及深度字段不自动补写。"
        "通过自动检查只说明提交包结构和规则约束成立，不替代人工核验。"
    )

    add_heading(doc, "十、阶段性结论", level=1)
    add_bullet(doc, "第十二周的56家公司队列已经转化为首批12家公司可运行结果，扩样从计划阶段进入实际处理阶段。")
    add_bullet(doc, "定位、候选、分类、复核和数据库披露形成一条完整链条，关键数字可以回到公司、页码和原始文本。")
    add_bullet(doc, "误命中没有被删除掩盖，而是单独记录26条排除原因，便于后续评估规则。")
    add_bullet(doc, "主体类型候选与备案编码/GP/LP深度字段分层处理，避免把名称线索当成确定事实。")
    add_bullet(doc, "当前仍是Auto阶段；未完成28条人工复核前，不宜计算或宣称最终准确率。")

    add_heading(doc, "十一、第十四周行动计划", level=1)
    add_table(
        doc,
        ["优先级", "任务", "具体操作", "验收产出"],
        [
            ["P0", "完成28条人工复核", "逐页核对事件真实性和核心字段，保留修改前后值", "review表完整率100%"],
            ["P0", "形成小规模Gold", "每家公司至少确认2条事件；合并跨页和重复披露", "不少于24条Gold事件"],
            ["P1", "量化Auto质量", "按事件识别和核心字段分别计算Precision、Recall、F1", "准确率报告及错误矩阵"],
            ["P1", "改进股权转让规则", "加入转让方/受让方/价款共现和公开转让负向规则", "误报率较本周下降"],
            ["P1", "补充基金深度核验", "仅对明确基金主体查询原文备案、GP和LP披露", "有证据则填，无证据保持空"],
            ["P2", "扩到下一批12家", "复用相同入口处理队列后续公司", "累计24家公司处理结果"],
        ],
        [1.8, 3.4, 7.3, 4.2],
        8.2,
    )

    doc.add_page_break()
    add_heading(doc, "十二、提交文件说明", level=1)
    add_table(
        doc,
        ["目录/文件", "作用"],
        [
            ["README.md", "统一运行命令、输入输出、统计口径和已知问题"],
            ["code/", "主流程、工作簿生成、报告生成、输出验证和统一入口"],
            ["data/source_texts/", "12家公司页码化文本输入副本"],
            ["data/derived/", "章节、候选、主体、交叉校验、复核队列等CSV"],
            ["database/", "PostgreSQL建表、导入、查询和结果披露"],
            ["outputs/霍泓锟_第十三周扩样核验结果.xlsx", "可筛选查看全部结果与人工填写入口"],
            ["report/霍泓锟_第十三周扩样抽取与人工核验报告.docx", "本周详实报告"],
            ["validation/", "自动检查和最终提交验证"],
            ["logs/", "主流程与数据库运行日志"],
        ],
        [7.0, 9.5],
        8.5,
    )

    add_heading(doc, "附录：关键统计口径", level=1)
    add_table(
        doc,
        ["名称", "定义", "不能解释为"],
        [
            ["Auto候选", "规则在页码化文本中发现的事件候选，包括保留和排除记录", "真实投资事件数"],
            ["保留候选", "未被明确负向规则排除、等待人工核验的候选", "Gold或Final"],
            ["高置信候选", "日期、金额、股数、比例等字段共现较充分", "一定正确"],
            ["旧页找回率", "本周定位页与第三周旧候选页的重合比例", "Precision、Recall或准确率"],
            ["主体类型候选", "根据名称与上下文得到的初步类型标签", "完成法律性质核验的最终分类"],
            ["未披露留空", "原文没有足够证据时不自动补值", "处理遗漏或程序失败"],
        ],
        [3.5, 8.0, 5.0],
        8.2,
    )

    doc.save(REPORT_PATH)
    print(REPORT_PATH)


if __name__ == "__main__":
    build_report()
